"""Pure Python review workbook export; never changes solved assignments."""
from collections import Counter
from datetime import date
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import scheduler_core as core

COLORS = {'D':'FFF2CC','E':'DDEBF7','N':'E2F0D9','OFF':'EEEEEE'}


def safe(value):
    """Prevent untrusted names/notes from becoming spreadsheet formulas."""
    if isinstance(value,str) and value.startswith(('=','+','-','@')):
        return "'"+value
    return value


def append(sheet,row):
    sheet.append([safe(v) for v in row])


def style(sheet,header_row=1):
    sheet.freeze_panes=f'C{header_row+1}'
    sheet.auto_filter.ref=f'A{header_row}:{get_column_letter(sheet.max_column)}{sheet.max_row}'
    for cell in sheet[header_row]:
        cell.fill=PatternFill('solid',fgColor='214F74')
        cell.font=Font(name='Microsoft JhengHei',color='FFFFFF',bold=True)
        cell.alignment=Alignment(wrap_text=True,vertical='center')
    sheet.row_dimensions[header_row].height=40
    for row in sheet.iter_rows(min_row=header_row+1):
        for cell in row:
            cell.font=Font(name='Microsoft JhengHei',size=10)
            cell.alignment=Alignment(wrap_text=True,vertical='center')
    for col in range(1,sheet.max_column+1):
        sheet.column_dimensions[get_column_letter(col)].width=18


def export_review(path,result,requests,report,ledger):
    if report['schedule_constraint_errors']:
        raise ValueError('排班未通過硬條件檢查，無法輸出班表')
    wb=Workbook(); intro=wb.active; intro.title='閱讀說明'
    for row in [ ['項目','內容'],['用途','供研究展示與人工審查，正式使用前須經院方核定。'],
        ['資料模式',report['mode']],['求解狀態',result.solver_status],
        ['輸入指紋',report['input_sha256']],['結果指紋',result.result_sha256],
        ['限制','四週假別須完整週期核定；不將 OFF 直接當超休；未確認不扣新增假。'],
        ['獨立判定','保留既有核心的年資或帶教累積判定，實際任用資格須由院方確認。'],
        ['最佳性','僅 OPTIMAL 表示目前設定之目標已達最佳；不代表所有相互衝突指標均各自最低。'],
        ['津貼設定','示範費率試算，非院方核定薪資或加班費。E：未滿15班每班500，滿15班每班700；N：未滿15班每班700，滿15班每班900。'],
        ['餘額核定','正式欄位依輸入核定旗標與程式檢查產生，並非系統向院方查證；空白表示資料不足或仍有警示，不代表零。'],
        ['隱私','內附示範為虛構人員。不得將真實名單、病歷或憑證隨研究作品對外公開。'] ]:
        append(intro,row)
    days=core.month_days(result.year,result.month); amap=result.assignment_map()
    leave={r['nurse_id']:r for r in ledger}
    annual={(r.nurse_id,r.day) for r in requests if r.kind==core.RequestKind.ANNUAL_LEAVE}
    for title,home,area in [('總班表',None,False),('總區域班表',None,True)]+[
        (s+'班班表',s,False) for s in 'DEN']+[(s+'區域班表',s,True) for s in 'DEN']:
        ws=wb.create_sheet(title)
        headers=['員編','姓名','原班別','到職日','D班數','E班數','N班數','跨班天數','目標H','工作H','加班H',
                 '特休期初H','補休期初H','特休試算H','補休試算H']+[d.day for d in days]
        append(ws,headers)
        members=[n for n in result.nurses if home is None or n.home_shift==home]
        for n in members:
            counts=core.shift_counts(result.assignments,n.nurse_id); b=leave.get(n.nurse_id,{})
            values=[n.nurse_id,n.name,n.home_shift,str(n.hire_date),counts['D'],counts['E'],counts['N'],
                core.cross_shift_days(result.assignments,{x.nurse_id:x for x in result.nurses},n.nurse_id),
                8*result.target_shifts[n.nurse_id],core.actual_hours(result.assignments,n.nurse_id),
                core.overtime_hours(result.assignments,result.target_shifts,n.nurse_id),
                b.get('annual_opening'),b.get('comp_opening'),b.get('annual_projected'),b.get('comp_projected')]
            for day in days:
                a=amap.get((n.nurse_id,day))
                if a:
                    cell=a.shift
                    if area: cell+='\n'+str(a.area or '未配置')
                    if a.preceptor_id: cell+='\n受教：'+a.preceptor_id
                    students=[x.nurse_id for x in result.assignments if x.day==day and x.preceptor_id==n.nurse_id]
                    if students: cell+='\n帶教：'+','.join(students)
                else: cell='特休' if (n.nurse_id,day) in annual else 'OFF'
                values.append(cell)
            append(ws,values)
        style(ws)
        for row in ws.iter_rows(min_row=2,min_col=16):
            for cell in row:
                code=str(cell.value).split('\n')[0]
                cell.fill=PatternFill('solid',fgColor=COLORS.get(code,'FCE4D6'))
                if '教：' in str(cell.value):
                    cell.font=Font(name='Microsoft JhengHei',size=10,bold=True,color='7030A0')
        for r in range(2,ws.max_row+1): ws.row_dimensions[r].height=55 if area else 35
        ws.column_dimensions['B'].width=24
    support=wb.create_sheet('跨班支援')
    append(support,['日期','員編','姓名','原班別','實際班別','區域'])
    profiles={n.nurse_id:n for n in result.nurses}
    for a in sorted(result.assignments,key=lambda a:(a.day,a.nurse_id)):
        n=profiles[a.nurse_id]
        if a.shift in 'DEN' and a.shift!=n.home_shift:
            append(support,[str(a.day),n.nurse_id,n.name,n.home_shift,a.shift,a.area])
    if support.max_row==1: append(support,['本月無跨班支援','','','','',''])
    allowance=wb.create_sheet('夜班津貼')
    append(allowance,['員編','姓名','E班數','E單班津貼','E合計','N班數','N單班津貼','N合計','總額'])
    for n in result.nurses:
        c=core.shift_counts(result.assignments,n.nurse_id); r=allowance.max_row+1
        append(allowance,[n.nurse_id,n.name,c['E'],None,None,c['N'],None,None,None])
        for column,formula in {'D':f'IF(C{r}<15,500,700)','E':f'C{r}*D{r}',
            'G':f'IF(F{r}<15,700,900)','H':f'F{r}*G{r}','I':f'E{r}+H{r}'}.items():
            allowance[f'{column}{r}']='='+formula
    balance=wb.create_sheet('休假時數結算'); questions=wb.create_sheet('特休補休詢問')
    append(balance,['員編','特休期初H','補休期初H','加班H','加班處理','特休試算H','補休試算H','特休正式H','補休正式H','提醒'])
    append(questions,['員編','本月需補假H','建議補休H','特休差額H','補休意願','已詢問特休','特休意願','提醒'])
    for b in ledger:
        append(balance,[b[k] for k in ['nurse_id','annual_opening','comp_opening','overtime_hours',
            'overtime_choice','annual_projected','comp_projected','annual_official','comp_official']]+['；'.join(b['warnings'])])
        append(questions,[b[k] for k in ['nurse_id','additional_leave_hours','suggested_comp_hours',
            'annual_shortfall_hours','comp_consent','annual_asked','annual_consent']]+['；'.join(b['warnings'])])
    if not ledger:
        append(balance,['尚未提供餘額，不以零代替']); append(questions,['請以 --leave-input 載入本人意願與已核定週期資料'])
    issues=wb.create_sheet('檢查與限制'); append(issues,['類型','代碼','說明'])
    for group in ['hard_errors','warnings']:
        for item in report[group]: append(issues,[group,item['code'],item['message']])
    append(issues,['限制','LOCAL_REVIEW','本次使用本機資料，未經 LIVE Google 來源驗證；來源檢查結果保留於報告。'])
    for ws in [intro,support,allowance,balance,questions,issues]: style(ws)
    intro.column_dimensions['B'].width=105; issues.column_dimensions['C'].width=105
    balance.column_dimensions['J'].width=80; questions.column_dimensions['H'].width=80
    wb.save(path)
    # Export round-trip regression: every original solved person/day cell is preserved.
    check=load_workbook(path,read_only=True,data_only=False)
    total=check['總班表']
    for row,n in enumerate(result.nurses,2):
        for column,day in enumerate(days,16):
            a=amap.get((n.nurse_id,day)); value=total.cell(row,column).value
            expected=a.shift if a else ('特休' if (n.nurse_id,day) in annual else 'OFF')
            assert str(value).split('\n')[0]==expected, (n.nurse_id,day)
    check.close()
